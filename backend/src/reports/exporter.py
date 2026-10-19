"""研究报告导出（PDF / Word）。"""

from __future__ import annotations

import re
from datetime import datetime
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Sequence, Tuple

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from backend.src.core.logger import get_logger
from backend.src.llm.client import SiliconFlowClient

logger = get_logger("reports")

EXPORT_DIR = Path(__file__).resolve().parents[3] / "outjson" / "reports"

_CJK_FONT_CANDIDATES: Tuple[Tuple[str, Path, int], ...] = (
    ("DrugIntelCJK", Path(r"C:\Windows\Fonts\simhei.ttf"), 0),
    ("DrugIntelCJK", Path(r"C:\Windows\Fonts\msyh.ttc"), 0),
    ("DrugIntelCJK", Path(r"C:\Windows\Fonts\simsun.ttc"), 0),
    ("DrugIntelCJK", Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"), 0),
    ("DrugIntelCJK", Path("/System/Library/Fonts/PingFang.ttc"), 0),
)


async def generate_report_summary(
    question: str,
    answer: str,
    sources: Sequence[Dict[str, Any]],
) -> str:
    llm = SiliconFlowClient.from_env()
    src = "\n".join(
        f"- {s.get('title')}" for s in sources[:8]
    )
    prompt = (
        f"请将以下问答整理为研究报告摘要（300字以内，中文）：\n"
        f"问题：{question}\n回答：{answer}\n来源：\n{src}"
    )
    try:
        return await llm.chat_completion(
            [
                {"role": "system", "content": "你是医药情报报告撰写助手。"},
                {"role": "user", "content": prompt},
            ]
        )
    except Exception as e:
        logger.warning("报告摘要生成失败: %s", e)
        return answer[:300]


@lru_cache(maxsize=1)
def _register_cjk_font() -> str | None:
    """注册系统中文字体；失败返回 None（英文仍可导出）。"""
    for name, path, subfont_index in _CJK_FONT_CANDIDATES:
        if not path.exists():
            continue
        try:
            if path.suffix.lower() == ".ttc":
                pdfmetrics.registerFont(
                    TTFont(name, str(path), subfontIndex=subfont_index)
                )
            else:
                pdfmetrics.registerFont(TTFont(name, str(path)))
            logger.info("PDF 中文字体已注册: %s (%s)", name, path)
            return name
        except Exception as e:
            logger.warning("注册字体失败 %s: %s", path, e)
    logger.warning("未找到可用中文字体，PDF 中文可能显示为方块")
    return None


def _escape_xml(text: str) -> str:
    return (
        (text or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def _markdown_to_reportlab_html(text: str) -> str:
    """
    将常见 Markdown 转为 reportlab Paragraph 可解析的简易 HTML。
    支持：**加粗**、*斜体*、换行、无序/有序列表行。
    """
    if not text:
        return ""

    lines: List[str] = []
    for raw in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw.rstrip()
        escaped = _escape_xml(line)

        # 无序/有序列表前缀保留可读形态
        m_ul = re.match(r"^(\s*)[-*+]\s+(.*)$", escaped)
        m_ol = re.match(r"^(\s*)(\d+)\.\s+(.*)$", escaped)
        if m_ul:
            indent, body = m_ul.group(1), m_ul.group(2)
            escaped = f"{indent}• {body}"
        elif m_ol:
            indent, num, body = m_ol.group(1), m_ol.group(2), m_ol.group(3)
            escaped = f"{indent}{num}. {body}"

        # **bold** / *italic*（先粗后斜，避免互相吞）
        escaped = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escaped)
        escaped = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"<i>\1</i>", escaped)
        lines.append(escaped if escaped else "&nbsp;")

    return "<br/>".join(lines)


def _pdf_styles():
    font_name = _register_cjk_font() or "Helvetica"
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "ReportTitle",
            parent=base["Title"],
            fontName=font_name,
            fontSize=18,
            leading=24,
            spaceAfter=8,
        ),
        "heading": ParagraphStyle(
            "ReportHeading",
            parent=base["Heading2"],
            fontName=font_name,
            fontSize=13,
            leading=18,
            spaceBefore=8,
            spaceAfter=4,
        ),
        "body": ParagraphStyle(
            "ReportBody",
            parent=base["BodyText"],
            fontName=font_name,
            fontSize=10.5,
            leading=16,
            spaceAfter=4,
        ),
        "meta": ParagraphStyle(
            "ReportMeta",
            parent=base["Normal"],
            fontName=font_name,
            fontSize=9,
            leading=13,
            textColor="#555555",
        ),
    }


def build_pdf_bytes(
    *,
    title: str,
    summary: str,
    question: str,
    answer: str,
    sources: Sequence[Dict[str, Any]],
) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm)
    styles = _pdf_styles()
    story = [
        Paragraph(_escape_xml(title), styles["title"]),
        Paragraph(
            _escape_xml(f"生成时间：{datetime.now().isoformat(timespec='seconds')}"),
            styles["meta"],
        ),
        Spacer(1, 0.35 * cm),
        Paragraph("摘要", styles["heading"]),
        Paragraph(_markdown_to_reportlab_html(summary), styles["body"]),
        Spacer(1, 0.2 * cm),
        Paragraph("问题", styles["heading"]),
        Paragraph(_markdown_to_reportlab_html(question), styles["body"]),
        Spacer(1, 0.2 * cm),
        Paragraph("回答", styles["heading"]),
        Paragraph(_markdown_to_reportlab_html(answer), styles["body"]),
        Spacer(1, 0.2 * cm),
        Paragraph("来源", styles["heading"]),
    ]
    for i, s in enumerate(sources, start=1):
        story.append(
            Paragraph(
                _markdown_to_reportlab_html(
                    f"[{i}] {s.get('title') or ''} {s.get('detail_url') or ''}"
                ),
                styles["body"],
            )
        )
    doc.build(story)
    return buf.getvalue()


def _add_markdown_paragraphs(document: Document, text: str) -> None:
    """Word：按行写入，并把 **加粗** 拆成 runs。"""
    if not text:
        document.add_paragraph("")
        return
    for raw in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw.rstrip()
        if not line:
            document.add_paragraph("")
            continue
        para = document.add_paragraph()
        # 列表前缀美化
        m_ul = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        m_ol = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m_ul:
            line = f"{m_ul.group(1)}• {m_ul.group(2)}"
        elif m_ol:
            line = f"{m_ol.group(1)}{m_ol.group(2)}. {m_ol.group(3)}"

        parts = re.split(r"(\*\*.+?\*\*)", line)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**") and len(part) >= 4:
                run = para.add_run(part[2:-2])
                run.bold = True
            else:
                para.add_run(part)


def build_docx_bytes(
    *,
    title: str,
    summary: str,
    question: str,
    answer: str,
    sources: Sequence[Dict[str, Any]],
) -> bytes:
    document = Document()
    document.add_heading(title, level=1)
    document.add_paragraph(f"生成时间：{datetime.now().isoformat(timespec='seconds')}")
    document.add_heading("摘要", level=2)
    _add_markdown_paragraphs(document, summary)
    document.add_heading("问题", level=2)
    _add_markdown_paragraphs(document, question)
    document.add_heading("回答", level=2)
    _add_markdown_paragraphs(document, answer)
    document.add_heading("来源", level=2)
    for i, s in enumerate(sources, start=1):
        document.add_paragraph(
            f"[{i}] {s.get('title') or ''} {s.get('detail_url') or ''}"
        )
    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()


async def export_report(
    *,
    question: str,
    answer: str,
    sources: Sequence[Dict[str, Any]],
    fmt: str = "pdf",
) -> Path:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    summary = await generate_report_summary(question, answer, sources)
    title = "DrugIntel 研究报告"
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    if fmt == "docx":
        data = build_docx_bytes(
            title=title, summary=summary, question=question, answer=answer, sources=sources
        )
        path = EXPORT_DIR / f"report_{stamp}.docx"
    else:
        data = build_pdf_bytes(
            title=title, summary=summary, question=question, answer=answer, sources=sources
        )
        path = EXPORT_DIR / f"report_{stamp}.pdf"
    path.write_bytes(data)
    return path
